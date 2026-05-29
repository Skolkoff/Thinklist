from __future__ import annotations

import shutil
import tempfile
import zipfile
from pathlib import Path
from textwrap import wrap

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "DOCS/practice/ФИО_Отчет_ИИАД_new - 2026.docx"
OUTPUT = ROOT / "DOCS/practice/Соколов_Никита_Евгеньевич_Отчет_преддипломная_практика_главы_1_3.docx"
ASSETS = ROOT / "DOCS/practice/report_assets/chapter1"


def font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    candidates = [
        "/System/Library/Fonts/Supplemental/Arial Bold.ttf" if bold else "/System/Library/Fonts/Supplemental/Arial.ttf",
        "/System/Library/Fonts/Supplemental/Arial Unicode.ttf",
        "/Library/Fonts/Arial.ttf",
        "/System/Library/Fonts/Supplemental/Times New Roman Bold.ttf" if bold else "/System/Library/Fonts/Supplemental/Times New Roman.ttf",
    ]
    for candidate in candidates:
        if candidate and Path(candidate).exists():
            return ImageFont.truetype(candidate, size)
    return ImageFont.load_default()


def draw_wrapped_text(draw: ImageDraw.ImageDraw, xy: tuple[int, int], text: str, width: int, fnt, fill: str) -> int:
    x, y = xy
    lines: list[str] = []
    for paragraph in text.split("\n"):
        current = ""
        for word in paragraph.split():
            trial = f"{current} {word}".strip()
            if draw.textbbox((0, 0), trial, font=fnt)[2] <= width:
                current = trial
            else:
                if current:
                    lines.append(current)
                current = word
        if current:
            lines.append(current)
    line_h = int(fnt.size * 1.25)
    for line in lines:
        draw.text((x, y), line, font=fnt, fill=fill)
        y += line_h
    return y


def make_process_diagram(path: Path) -> None:
    img = Image.new("RGB", (1800, 720), "#ffffff")
    draw = ImageDraw.Draw(img)
    title = font(42, True)
    body = font(27)
    small = font(23)
    draw.text((70, 44), "Бизнес-процесс работы с научными публикациями", font=title, fill="#111827")

    steps = [
        ("1", "Поиск и импорт", "arXiv, демо-данные,\nметаданные статьи"),
        ("2", "Нормализация", "теги, авторы,\nгод, источник"),
        ("3", "Просмотр ленты", "поиск, фильтры,\nкарточка статьи"),
        ("4", "Пользовательские сигналы", "сохранение, лайк,\nкомментарий, скрытие"),
        ("5", "Гибридное ранжирование", "TF-IDF, SVD,\nMLP-реранкер"),
        ("6", "Результат", "персональная выдача\nи коллекции"),
    ]
    x0, y0 = 70, 180
    box_w, box_h, gap = 255, 250, 28
    for i, (num, heading, desc) in enumerate(steps):
        x = x0 + i * (box_w + gap)
        draw.rounded_rectangle((x, y0, x + box_w, y0 + box_h), radius=28, fill="#eff6ff", outline="#2563eb", width=4)
        draw.ellipse((x + 22, y0 + 22, x + 78, y0 + 78), fill="#2563eb")
        draw.text((x + 43, y0 + 32), num, font=font(26, True), fill="#ffffff", anchor="mm")
        draw_wrapped_text(draw, (x + 24, y0 + 92), heading, box_w - 48, body, "#111827")
        draw_wrapped_text(draw, (x + 24, y0 + 154), desc, box_w - 48, small, "#374151")
        if i < len(steps) - 1:
            ax = x + box_w + 8
            ay = y0 + box_h // 2
            draw.line((ax, ay, ax + gap - 16, ay), fill="#64748b", width=5)
            draw.polygon([(ax + gap - 16, ay), (ax + gap - 34, ay - 13), (ax + gap - 34, ay + 13)], fill="#64748b")

    draw.rounded_rectangle((70, 520, 1730, 650), radius=24, fill="#f8fafc", outline="#cbd5e1", width=3)
    draw_wrapped_text(
        draw,
        (104, 548),
        "Ключевая идея процесса: каждый контакт пользователя со статьей превращается в слабый сигнал, "
        "который уточняет последующую выдачу и помогает поддерживать баланс между точностью, новизной и смежными направлениями.",
        1600,
        font(27),
        "#1f2937",
    )
    img.save(path)


def make_architecture_diagram(path: Path) -> None:
    img = Image.new("RGB", (1800, 980), "#ffffff")
    draw = ImageDraw.Draw(img)
    title = font(42, True)
    body = font(27)
    small = font(24)
    draw.text((70, 44), "Архитектура информационной системы SciVise", font=title, fill="#111827")

    boxes = {
        "frontend": (90, 190, 530, 390, "#ecfeff", "#0891b2", "Frontend", "React 19, TypeScript,\nVite, Tailwind CSS"),
        "backend": (680, 190, 1120, 390, "#f0fdf4", "#16a34a", "Backend API", "FastAPI, SQLAlchemy,\nJWT, бизнес-логика"),
        "ml": (1270, 190, 1710, 390, "#fefce8", "#ca8a04", "ML-service", "гибридный скоринг,\nсемантика, rerank"),
        "db": (680, 565, 1120, 765, "#f8fafc", "#475569", "PostgreSQL", "пользователи, статьи,\nтеги, события, коллекции"),
        "source": (1270, 565, 1710, 765, "#fff7ed", "#ea580c", "Внешние источники", "arXiv и импорт\nметаданных статей"),
    }
    for x1, y1, x2, y2, fill, outline, heading, desc in boxes.values():
        draw.rounded_rectangle((x1, y1, x2, y2), radius=28, fill=fill, outline=outline, width=4)
        draw_wrapped_text(draw, (x1 + 34, y1 + 36), heading, x2 - x1 - 68, body, "#111827")
        draw_wrapped_text(draw, (x1 + 34, y1 + 96), desc, x2 - x1 - 68, small, "#374151")

    def arrow(start: tuple[int, int], end: tuple[int, int], label: str) -> None:
        draw.line((*start, *end), fill="#334155", width=5)
        ex, ey = end
        sx, sy = start
        if abs(ex - sx) >= abs(ey - sy):
            sign = 1 if ex > sx else -1
            draw.polygon([(ex, ey), (ex - sign * 22, ey - 14), (ex - sign * 22, ey + 14)], fill="#334155")
        else:
            sign = 1 if ey > sy else -1
            draw.polygon([(ex, ey), (ex - 14, ey - sign * 22), (ex + 14, ey - sign * 22)], fill="#334155")
        lx, ly = (sx + ex) // 2, (sy + ey) // 2
        draw.rounded_rectangle((lx - 105, ly - 24, lx + 105, ly + 24), radius=16, fill="#ffffff", outline="#cbd5e1", width=2)
        draw.text((lx, ly), label, font=font(20), fill="#334155", anchor="mm")

    arrow((530, 290), (680, 290), "HTTP API")
    arrow((1120, 290), (1270, 290), "recommend")
    arrow((900, 390), (900, 565), "SQL")
    arrow((1490, 565), (1490, 390), "metadata")
    arrow((1270, 665), (1120, 665), "ingest")

    draw.rounded_rectangle((90, 840, 1710, 920), radius=22, fill="#f1f5f9", outline="#cbd5e1", width=3)
    draw.text((118, 880), "Развертывание: Docker Compose, локальные порты 5740 / 8000 / 8001 / 5432.", font=small, fill="#1f2937")
    img.save(path)


def make_component_chart(path: Path) -> None:
    data = [
        ("UI-модули", 11, "#2563eb"),
        ("API-группы", 8, "#16a34a"),
        ("DB-сущности", 18, "#64748b"),
        ("ML-факторы", 7, "#ca8a04"),
        ("QA-сценарии", 12, "#dc2626"),
    ]
    img = Image.new("RGB", (1500, 850), "#ffffff")
    draw = ImageDraw.Draw(img)
    draw.text((70, 44), "Состав MVP по ключевым подсистемам", font=font(42, True), fill="#111827")
    left, top, right, bottom = 180, 170, 1360, 670
    draw.line((left, bottom, right, bottom), fill="#334155", width=4)
    draw.line((left, top, left, bottom), fill="#334155", width=4)
    max_v = max(v for _, v, _ in data)
    bar_w, gap = 150, 75
    for i, (label, value, color) in enumerate(data):
        x = left + 80 + i * (bar_w + gap)
        h = int((value / max_v) * 410)
        y = bottom - h
        draw.rounded_rectangle((x, y, x + bar_w, bottom), radius=18, fill=color)
        draw.text((x + bar_w // 2, y - 34), str(value), font=font(30, True), fill="#111827", anchor="mm")
        words = "\n".join(wrap(label, 10))
        draw.multiline_text((x + bar_w // 2, bottom + 28), words, font=font(24), fill="#374151", anchor="ma", align="center", spacing=5)
    draw.rounded_rectangle((110, 740, 1390, 805), radius=18, fill="#f8fafc", outline="#cbd5e1", width=2)
    draw.text((140, 773), "Оценка отражает текущий учебный MVP: реализованный интерфейс, API, БД, ML-скоринг и сценарии проверки.", font=font(23), fill="#374151")
    img.save(path)


def make_recommender_pipeline(path: Path) -> None:
    img = Image.new("RGB", (1800, 820), "#ffffff")
    draw = ImageDraw.Draw(img)
    draw.text((70, 42), "Пайплайн гибридной рекомендательной системы", font=font(42, True), fill="#111827")
    stages = [
        ("Источники", "публикации,\nпрофиль,\nдействия"),
        ("Нормализация", "текст,\nтеги,\nметаданные"),
        ("Признаки", "TF-IDF,\nSVD,\nинтересы"),
        ("Кандидаты", "контентная\nпохожесть"),
        ("Скоринг", "гибридная\nформула"),
        ("Реранжирование", "нейронная\nмодель"),
        ("Выдача", "лента и\nколлекции"),
    ]
    x0, y0 = 70, 180
    box_w, box_h, gap = 218, 220, 30
    for i, (title, desc) in enumerate(stages):
        x = x0 + i * (box_w + gap)
        fill = ["#eef2ff", "#ecfeff", "#f0fdf4", "#fff7ed", "#fefce8", "#fdf2f8", "#f8fafc"][i]
        outline = ["#4f46e5", "#0891b2", "#16a34a", "#ea580c", "#ca8a04", "#db2777", "#475569"][i]
        draw.rounded_rectangle((x, y0, x + box_w, y0 + box_h), radius=24, fill=fill, outline=outline, width=4)
        draw_wrapped_text(draw, (x + 22, y0 + 34), title, box_w - 44, font(28, True), "#111827")
        draw_wrapped_text(draw, (x + 22, y0 + 104), desc, box_w - 44, font(25), "#374151")
        if i < len(stages) - 1:
            ax, ay = x + box_w + 8, y0 + box_h // 2
            draw.line((ax, ay, ax + gap - 16, ay), fill="#64748b", width=5)
            draw.polygon([(ax + gap - 16, ay), (ax + gap - 34, ay - 13), (ax + gap - 34, ay + 13)], fill="#64748b")
    draw.rounded_rectangle((120, 520, 1680, 705), radius=28, fill="#f8fafc", outline="#cbd5e1", width=3)
    draw_wrapped_text(
        draw,
        (160, 552),
        "Пайплайн объединяет контентные признаки статьи, пользовательские интересы и поведенческие сигналы. "
        "Финальная выдача формируется не только по близости темы, но и с учетом свежести, разнообразия и смежных направлений.",
        1480,
        font(28),
        "#1f2937",
    )
    img.save(path)


def make_scoring_chart(path: Path) -> None:
    data = [
        ("semantic", 35, "#2563eb"),
        ("tf-idf", 25, "#0891b2"),
        ("interest", 20, "#16a34a"),
        ("behavior", 10, "#ca8a04"),
        ("freshness", 5, "#ea580c"),
        ("diversity", 5, "#db2777"),
    ]
    img = Image.new("RGB", (1500, 780), "#ffffff")
    draw = ImageDraw.Draw(img)
    draw.text((70, 42), "Компоненты гибридного скоринга", font=font(42, True), fill="#111827")
    left, top, bar_h = 220, 170, 72
    scale = 900 / 35
    for i, (label, value, color) in enumerate(data):
        y = top + i * 88
        draw.text((70, y + 36), label, font=font(26, True), fill="#111827", anchor="lm")
        draw.rounded_rectangle((left, y, left + 900, y + bar_h), radius=22, fill="#f1f5f9")
        draw.rounded_rectangle((left, y, left + int(value * scale), y + bar_h), radius=22, fill=color)
        draw.text((left + int(value * scale) + 24, y + 36), f"{value}%", font=font(27, True), fill="#111827", anchor="lm")
    draw.rounded_rectangle((70, 700, 1430, 748), radius=16, fill="#f8fafc", outline="#cbd5e1", width=2)
    draw.text((96, 724), "Diversity указан как корректирующий штраф за однообразие выдачи, а не как самостоятельное увеличение оценки.", font=font(22), fill="#374151", anchor="lm")
    img.save(path)


def make_neural_reranker_diagram(path: Path) -> None:
    img = Image.new("RGB", (1600, 880), "#ffffff")
    draw = ImageDraw.Draw(img)
    draw.text((70, 42), "Нейронный реранкер в рекомендательном пайплайне", font=font(42, True), fill="#111827")
    layers = [
        ("Входные признаки", 5, 130, "#e0f2fe", "#0284c7"),
        ("Скрытый слой 1", 4, 430, "#dcfce7", "#16a34a"),
        ("Скрытый слой 2", 3, 730, "#fef3c7", "#ca8a04"),
        ("Оценка статьи", 1, 1030, "#fce7f3", "#db2777"),
    ]
    positions: list[list[tuple[int, int]]] = []
    for title, count, x, fill, outline in layers:
        draw.text((x + 85, 150), title, font=font(25, True), fill="#111827", anchor="mm")
        ys = [250 + i * 95 for i in range(count)]
        if count == 1:
            ys = [420]
        layer_pos = []
        for y in ys:
            draw.ellipse((x, y, x + 170, y + 70), fill=fill, outline=outline, width=4)
            layer_pos.append((x + 170, y + 35))
            draw.ellipse((x + 68, y + 18, x + 102, y + 52), fill=outline)
        positions.append(layer_pos)
    for left_layer, right_layer in zip(positions, positions[1:]):
        for sx, sy in left_layer:
            for ex, ey in right_layer:
                draw.line((sx, sy, ex - 170, ey), fill="#cbd5e1", width=2)
    draw.rounded_rectangle((70, 720, 1530, 820), radius=22, fill="#f8fafc", outline="#cbd5e1", width=3)
    draw_wrapped_text(
        draw,
        (100, 748),
        "Реранкер получает признаки пары «пользователь — статья» и уточняет порядок кандидатов после базового контентного отбора. "
        "В MVP его задача — показать применимость нейросетевого слоя без усложнения инфраструктуры.",
        1380,
        font(25),
        "#374151",
    )
    img.save(path)


def make_adjacent_domains_diagram(path: Path) -> None:
    img = Image.new("RGB", (1500, 900), "#ffffff")
    draw = ImageDraw.Draw(img)
    draw.text((70, 42), "Подмешивание публикаций из смежных направлений", font=font(42, True), fill="#111827")
    center = (750, 430)
    domains = [
        ("Computer\nScience", center, 170, "#2563eb"),
        ("Mathematics", (420, 250), 120, "#16a34a"),
        ("Physics", (1050, 250), 120, "#ca8a04"),
        ("Biology", (420, 620), 120, "#db2777"),
        ("Social\nSciences", (1050, 620), 120, "#0891b2"),
    ]
    for label, (x, y), r, color in domains[1:]:
        draw.line((center[0], center[1], x, y), fill="#cbd5e1", width=8)
    for label, (x, y), r, color in domains:
        draw.ellipse((x - r, y - r, x + r, y + r), fill=color, outline="#111827", width=4)
        draw.multiline_text((x, y), label, font=font(29, True), fill="#ffffff", anchor="mm", align="center", spacing=6)
    draw.rounded_rectangle((165, 780, 1335, 835), radius=18, fill="#f8fafc", outline="#cbd5e1", width=2)
    draw.text((750, 807), "Базовая пропорция выдачи: около 80% профильной темы и 20% смежных публикаций.", font=font(24), fill="#374151", anchor="mm")
    img.save(path)


def make_frontend_flow_diagram(path: Path) -> None:
    img = Image.new("RGB", (1700, 900), "#ffffff")
    draw = ImageDraw.Draw(img)
    draw.text((70, 42), "Логика пользовательских переходов в интерфейсе SciVise", font=font(42, True), fill="#111827")
    nodes = {
        "Гость": (110, 180, 390, 300, "#f8fafc", "#475569"),
        "Регистрация": (520, 180, 850, 300, "#eef2ff", "#4f46e5"),
        "Онбординг": (980, 180, 1310, 300, "#ecfeff", "#0891b2"),
        "Лента": (520, 420, 850, 540, "#f0fdf4", "#16a34a"),
        "Карточка статьи": (980, 420, 1310, 540, "#fff7ed", "#ea580c"),
        "Коллекции": (110, 650, 390, 770, "#fefce8", "#ca8a04"),
        "Профиль": (520, 650, 850, 770, "#fdf2f8", "#db2777"),
        "Настройки": (980, 650, 1310, 770, "#f1f5f9", "#334155"),
    }
    for label, (x1, y1, x2, y2, fill, outline) in nodes.items():
        draw.rounded_rectangle((x1, y1, x2, y2), radius=26, fill=fill, outline=outline, width=4)
        draw.text(((x1 + x2) // 2, (y1 + y2) // 2), label, font=font(29, True), fill="#111827", anchor="mm")

    def arrow(a: str, b: str, label: str = "") -> None:
        ax1, ay1, ax2, ay2, *_ = nodes[a]
        bx1, by1, bx2, by2, *_ = nodes[b]
        start = (ax2, (ay1 + ay2) // 2) if ax2 < bx1 else ((ax1 + ax2) // 2, ay2)
        end = (bx1, (by1 + by2) // 2) if ax2 < bx1 else ((bx1 + bx2) // 2, by1)
        draw.line((*start, *end), fill="#64748b", width=5)
        ex, ey = end
        draw.polygon([(ex, ey), (ex - 18, ey - 12), (ex - 18, ey + 12)], fill="#64748b")
        if label:
            lx, ly = (start[0] + end[0]) // 2, (start[1] + end[1]) // 2
            draw.rounded_rectangle((lx - 95, ly - 22, lx + 95, ly + 22), radius=14, fill="#ffffff", outline="#cbd5e1", width=2)
            draw.text((lx, ly), label, font=font(18), fill="#334155", anchor="mm")

    arrow("Гость", "Регистрация")
    arrow("Регистрация", "Онбординг")
    arrow("Онбординг", "Лента")
    arrow("Лента", "Карточка статьи", "открыть")
    arrow("Лента", "Коллекции")
    arrow("Лента", "Профиль")
    arrow("Профиль", "Настройки")
    arrow("Карточка статьи", "Коллекции", "сохранить")
    draw.rounded_rectangle((1370, 360, 1605, 590), radius=24, fill="#f8fafc", outline="#cbd5e1", width=3)
    draw_wrapped_text(draw, (1400, 395), "Ключевой принцип: пользователь может вернуться к ленте из любого сценария без потери контекста.", 175, font(23), "#374151")
    img.save(path)


def make_title_page_image(base_image: bytes, path: Path) -> None:
    img = Image.open(__import__("io").BytesIO(base_image)).convert("RGB")
    draw = ImageDraw.Draw(img)
    regular = font(12)
    small = font(8)
    bottom = font(14)

    def cover(box: tuple[int, int, int, int]) -> None:
        draw.rectangle(box, fill="#ffffff")

    def center_text(box: tuple[int, int, int, int], text: str, fnt, fill: str = "#111111") -> None:
        x1, y1, x2, y2 = box
        draw.text(((x1 + x2) // 2, (y1 + y2) // 2), text, font=fnt, fill=fill, anchor="mm")

    cover((248, 588, 512, 633))
    center_text((248, 592, 512, 616), "Соколова Никиты Евгеньевича", regular)

    cover((235, 640, 500, 668))
    center_text((235, 640, 500, 668), "кафедра информационных систем", regular)

    cover((240, 682, 365, 708))
    center_text((240, 682, 365, 708), "13.04.2026 г.", regular)
    cover((428, 682, 555, 708))
    center_text((428, 682, 555, 708), "10.05.2026 г.", regular)

    cover((305, 815, 430, 850))
    center_text((305, 815, 430, 850), "г. Москва 2026 г.", bottom)

    # Restore small under-line labels that overlap patched areas.
    center_text((360, 620, 470, 633), "(фамилия, имя, отчество)", small)
    img.save(path)


def replace_docx_media(docx_path: Path, media_name: str, replacement: Path) -> None:
    with tempfile.NamedTemporaryFile(delete=False, suffix=".docx") as tmp:
        tmp_path = Path(tmp.name)
    with zipfile.ZipFile(docx_path, "r") as zin, zipfile.ZipFile(tmp_path, "w", zipfile.ZIP_DEFLATED) as zout:
        replacement_bytes = replacement.read_bytes()
        for item in zin.infolist():
            data = replacement_bytes if item.filename == media_name else zin.read(item.filename)
            zout.writestr(item, data)
    shutil.move(str(tmp_path), docx_path)


def set_paragraph_font(paragraph, size: int = 14, bold: bool | None = None) -> None:
    for run in paragraph.runs:
        run.font.name = "Times New Roman"
        run.font.size = Pt(size)
        if bold is not None:
            run.bold = bold


def add_caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph(text)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_font(p, 12)
    for r in p.runs:
        r.italic = True


def add_figure(doc: Document, image_path: Path, width: float, caption: str, description: str) -> None:
    doc.add_picture(str(image_path), width=Inches(width))
    add_caption(doc, caption)
    add_body(doc, description)


def add_body(doc: Document, text: str) -> None:
    p = doc.add_paragraph(text)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Inches(0.49)
    p.paragraph_format.line_spacing = 1.5
    set_paragraph_font(p)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph()
        p.add_run(f"— {item}")
        p.paragraph_format.line_spacing = 1.5
        set_paragraph_font(p)


def add_numbered(doc: Document, items: list[str]) -> None:
    for index, item in enumerate(items, start=1):
        p = doc.add_paragraph()
        p.add_run(f"{index}. {item}")
        p.paragraph_format.line_spacing = 1.5
        set_paragraph_font(p)


def add_heading(doc: Document, level: int, text: str):
    p = doc.add_heading(text, level=level)
    set_paragraph_font(p, 14 if level > 1 else 16, True)
    return p


def fill_table(table, rows: list[list[str]]) -> None:
    table.style = "Table Grid"
    for r, row in enumerate(rows):
        cells = table.rows[r].cells
        for c, value in enumerate(row):
            cells[c].text = value
            for paragraph in cells[c].paragraphs:
                set_paragraph_font(paragraph, 11, bold=(r == 0))


def add_table(doc: Document, caption: str, rows: list[list[str]]) -> None:
    add_caption(doc, caption)
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    fill_table(table, rows)


def remove_range_between(doc: Document, start_text: str, end_text: str):
    start = None
    end = None
    for p in doc.paragraphs:
        is_heading_1 = p.style is not None and p.style.name.startswith("Heading 1")
        if start is None and is_heading_1 and start_text in p.text:
            start = p._p
        elif start is not None and is_heading_1 and end_text in p.text:
            end = p._p
            break
    if start is None or end is None:
        raise RuntimeError("Could not find chapter boundaries in the template")
    body = doc._body._element
    children = list(body)
    start_idx = children.index(start)
    end_idx = children.index(end)
    for el in children[start_idx:end_idx]:
        body.remove(el)
    return end


def move_new_body_elements_before(doc: Document, marker, original_len: int, section_properties) -> None:
    body = doc._body._element
    children = list(body)
    section_index = children.index(section_properties)
    new_elems = children[original_len - 1:section_index]
    for el in new_elems:
        body.remove(el)
        marker.addprevious(el)


def replace_intro(doc: Document) -> None:
    marker = remove_range_between(doc, "Введение", "Анализ бизнес-процесса / ов")
    body = doc._body._element
    original_len = len(list(body))
    section_properties = list(body)[-1]

    add_heading(doc, 1, "Введение")
    intro_paragraphs = [
        "Современная научная коммуникация характеризуется быстрым ростом числа публикаций, препринтов и междисциплинарных исследований. Пользователю становится все сложнее самостоятельно отбирать релевантные материалы, отслеживать новые статьи и находить работы из смежных областей.",
        "Особенно заметна эта проблема в области искусственного интеллекта и анализа данных. Новые методы появляются на стыке компьютерных наук, математики, физики, биологии и социальных наук, поэтому классический поиск по ключевым словам часто оказывается недостаточным.",
        "Актуальность работы заключается в необходимости разработки информационной системы, которая помогает пользователю ориентироваться в массиве научных публикаций и получать персонализированные рекомендации. Такая система должна учитывать содержание статей, интересы пользователя и историю его взаимодействия с материалами.",
        "В рамках преддипломной практики рассматривается проект SciVise — веб-приложение для поиска, просмотра, сохранения и рекомендации научных публикаций. Проект предназначен для поддержки учебной и исследовательской деятельности студентов, преподавателей и молодых исследователей.",
        "Цель проекта — разработать прототип SciVise, который формирует персональную ленту научных публикаций, сохраняет материалы в коллекции и объяснимо ранжирует статьи на основе содержания, интересов пользователя и его действий.",
        "Цель преддипломной практики — подготовить проектную основу ВКР: описать предметную область, сформулировать требования, спроектировать архитектуру, определить состав рекомендательного пайплайна и подготовить систему к экспериментальной проверке.",
        "Для достижения цели поставлены конкретные задачи: исследовать пользовательский процесс, выбрать стек технологий, описать модель данных, реализовать базовые сценарии, спроектировать гибридное ранжирование и определить критерии проверки качества рекомендаций.",
        "Объектом исследования выступает процесс поиска и отбора научных публикаций пользователем. Предметом исследования являются методы персонализации выдачи, гибридного ранжирования и учета пользовательских сигналов при формировании рекомендаций.",
        "Для достижения цели необходимо проанализировать бизнес-процесс работы с научными публикациями, определить требования пользователей, описать архитектуру системы, сформировать техническое задание и подготовить основу для последующей реализации рекомендательного модуля.",
        "В ходе работы применяются методы анализа предметной области, проектирования информационных систем, моделирования пользовательских сценариев, анализа требований, программной реализации веб-приложений и экспериментальной проверки рекомендательных алгоритмов.",
        "Практическая значимость проекта состоит в создании прототипа системы, который объединяет пользовательский интерфейс, серверную логику, базу данных и ML-сервис. Такой прототип позволяет демонстрировать полный цикл персонализированной работы с научными публикациями.",
        "Структура отчета соответствует логике преддипломной практики. В первой главе рассматриваются предметная область, требования и техническое задание. Последующие главы предназначены для описания архитектуры, реализации, тестирования и оценки полученных результатов.",
    ]
    for text in intro_paragraphs:
        add_body(doc, text)

    move_new_body_elements_before(doc, marker, original_len, section_properties)


def replace_chapter2(doc: Document, assets: dict[str, Path]) -> None:
    marker = remove_range_between(doc, "проектирование АРХИТЕКТУРЫ", "Frontend-разработка / модификация")
    body = doc._body._element
    original_len = len(list(body))
    section_properties = list(body)[-1]

    h1 = add_heading(doc, 1, "2 Проектирование архитектуры рекомендательной системы SciVise")
    h1.paragraph_format.page_break_before = True

    add_heading(doc, 2, "2.1 Разработка предложений по архитектуре рекомендательного модуля")
    add_body(doc, "Рекомендательный модуль SciVise проектируется как отдельный вычислительный контур, связанный с основным приложением через API. Такое разделение позволяет развивать алгоритмы ранжирования независимо от пользовательского интерфейса и основной серверной логики.")
    add_body(doc, "Архитектура строится вокруг трех типов данных: метаданные научных публикаций, текстовые признаки статей и пользовательские сигналы. Метаданные помогают фильтровать материалы, текстовые признаки отражают смысл публикации, а сигналы пользователя уточняют персональные предпочтения.")
    add_body(doc, "Базовым подходом выбран гибридный пайплайн. Он сочетает контентную близость, интересы пользователя, поведенческие события, свежесть публикации, нейронный реранкер и механизм разнообразия выдачи. Такой подход лучше подходит для учебного MVP, чем полностью нейросетевая система.")
    add_figure(
        doc,
        assets["pipeline"],
        6.6,
        "Рисунок 5 — Пайплайн гибридной рекомендательной системы SciVise",
        "На рисунке 5 показан полный путь данных в рекомендательном модуле: от исходных публикаций, профиля и действий пользователя до нормализации, формирования признаков, скоринга, нейронного реранжирования и выдачи персональной ленты.",
    )

    add_heading(doc, 2, "2.2 Архитектурная спецификация компонентов рекомендательной системы")
    add_body(doc, "В архитектуре выделяются четыре основных компонента: клиентское приложение, Backend API, база данных и ML-service. Клиентское приложение отображает ленту и фиксирует действия пользователя. Backend API валидирует запросы и обращается к базе данных. ML-service рассчитывает персонализированную выдачу.")
    add_body(doc, "База данных хранит публикации, теги, пользователей, интересы, коллекции, комментарии и события взаимодействия. ML-service получает данные из базы, строит признаки, вычисляет кандидатов и возвращает отсортированный список публикаций. Такой обмен обеспечивает понятный контур ответственности каждого компонента.")
    add_table(
        doc,
        "Таблица 4 — Данные и признаки рекомендательной системы",
        [
            ["Группа данных", "Примеры", "Использование в рекомендациях"],
            ["Метаданные статьи", "название, авторы, год, источник, arXiv-категории", "фильтрация, отображение карточки, первичное ранжирование"],
            ["Текстовые признаки", "название, аннотация, тематические теги", "TF-IDF, семантическое представление, контентная похожесть"],
            ["Профиль пользователя", "интересы, язык интерфейса, сохраненные материалы", "персонализация и усиление релевантных тем"],
            ["Поведенческие события", "лайк, сохранение, скрытие, просмотр, комментарий", "обновление предпочтений и расчет behavioral score"],
            ["Контекст выдачи", "выбранная тема, поиск, свежесть, разнообразие", "коррекция итогового ранга и подмешивание смежных направлений"],
        ],
    )

    add_heading(doc, 2, "2.3 Разработка прототипа модели рекомендаций")
    add_body(doc, "Прототип модели рекомендаций состоит из двух уровней. Первый уровень формирует кандидатов на основе контентной близости между публикациями и интересами пользователя. Второй уровень уточняет порядок кандидатов с помощью гибридной формулы и нейронного реранкера.")
    add_body(doc, "Гибридный скоринг выбран потому, что он объясним и устойчив на небольшом наборе данных. В отличие от сложной end-to-end нейросети, такая модель позволяет явно контролировать вклад смысловой близости, интересов, поведения, свежести и разнообразия.")
    add_table(
        doc,
        "Таблица 5 — Компоненты итоговой оценки публикации",
        [
            ["Компонент", "Назначение", "Ожидаемый эффект"],
            ["Semantic score", "сравнивает плотные представления публикаций", "повышает статьи, близкие по смыслу"],
            ["TF-IDF score", "учитывает совпадение терминов в названии и аннотации", "поддерживает точность по ключевым понятиям"],
            ["Interest score", "сопоставляет статью с интересами пользователя", "делает ленту персональной"],
            ["Behavior score", "использует сохранения, реакции и скрытия", "адаптирует выдачу к фактическим действиям"],
            ["Freshness score", "учитывает год публикации", "поддерживает появление новых материалов"],
            ["Diversity penalty", "снижает однообразие близких статей", "добавляет смежные направления"],
        ],
    )
    add_figure(
        doc,
        assets["scoring"],
        6.4,
        "Рисунок 6 — Весовые компоненты гибридного скоринга",
        "На рисунке 6 отражено соотношение факторов итоговой оценки публикации. Наибольший вклад вносит семантическая близость, далее учитываются TF-IDF-признаки, интересы пользователя, поведенческие события, свежесть и штраф за однообразие выдачи.",
    )

    add_heading(doc, 3, "2.3.1 Формирование набора данных и признаков")
    add_body(doc, "Набор данных формируется из публикаций, тематических тегов и пользовательских событий. Для каждой статьи используются название, аннотация, год публикации, источник и тематические категории. Для пользователя используются интересы, сохраненные статьи и события взаимодействия с карточками.")
    add_body(doc, "На этапе признакового описания текст статьи преобразуется в TF-IDF-вектор. Затем с помощью понижения размерности формируется более компактное семантическое представление. Это снижает шум разреженных признаков и делает сравнение публикаций более устойчивым.")
    add_figure(
        doc,
        assets["adjacent"],
        6.2,
        "Рисунок 7 — Схема подмешивания публикаций из смежных направлений",
        "На рисунке 7 показан принцип междисциплинарного расширения выдачи. Основная часть рекомендаций остается в профильной области, но небольшая доля материалов поступает из связанных направлений, что помогает находить новые исследовательские идеи.",
    )

    add_heading(doc, 3, "2.3.2 Выбор модели и инструментальных средств")
    add_body(doc, "Для MVP выбран стек, который можно запустить локально без GPU: TF-IDF для текстовых признаков, TruncatedSVD для плотного представления, MLPRegressor для нейронного реранжирования и FastAPI для выделенного ML-сервиса. Такой набор достаточен для демонстрации дипломной идеи.")
    add_body(doc, "Нейронная сеть используется не как единственный источник рекомендаций, а как дополнительный слой уточнения порядка кандидатов. Это снижает риск нестабильной выдачи и сохраняет возможность объяснить, почему статья получила высокий итоговый ранг.")
    add_figure(
        doc,
        assets["neural"],
        6.2,
        "Рисунок 8 — Упрощенная архитектура нейронного реранкера",
        "На рисунке 8 представлена роль нейронного слоя в системе. Реранкер получает признаки пары «пользователь — статья» и уточняет порядок уже отобранных кандидатов, не заменяя базовую гибридную модель целиком.",
    )

    add_heading(doc, 2, "2.4 Обучение и расчет рекомендаций")
    add_body(doc, "Процесс обучения в MVP можно рассматривать как периодическое обновление признаков и параметров реранкера. При появлении новых публикаций пересчитываются текстовые представления. При накоплении пользовательских событий уточняется модель поведения и веса персонализации.")
    add_body(doc, "Расчет рекомендаций выполняется в несколько шагов: сбор данных пользователя, формирование кандидатного набора, расчет контентных оценок, применение гибридной формулы, нейронное уточнение порядка и финальная коррекция разнообразия. Результат возвращается в ленту пользователя.")
    add_figure(
        doc,
        assets["article_modal"],
        6.4,
        "Рисунок 9 — Карточка статьи как источник пользовательских сигналов",
        "На рисунке 9 показана карточка публикации, через которую пользователь изучает аннотацию, сохраняет материал и оставляет комментарий. Эти действия становятся источником слабых сигналов для дальнейшей персонализации рекомендаций.",
    )

    add_heading(doc, 2, "2.5 Разработка программного модуля рекомендательной системы")
    add_body(doc, "Программный модуль рекомендаций размещается в ML-service и взаимодействует с Backend API. Backend API отвечает за пользовательские запросы и безопасность, а ML-service концентрирует вычислительную часть: подготовку признаков, расчет score-компонентов и формирование итогового списка.")
    add_body(doc, "Такое разделение полезно для дальнейших экспериментов. Внутри ML-service можно менять алгоритм ранжирования, добавлять новые признаки и проверять метрики без изменения интерфейса. Для ВКР это позволяет отдельно описать программную реализацию и экспериментальную часть.")
    add_body(doc, "Важной особенностью модуля является поддержка смежных направлений. Для научного поиска это критично, потому что пользователь может находить полезные идеи в соседних областях. Поэтому выдача должна сочетать профильные статьи и ограниченную долю материалов из связанных тематик.")

    add_heading(doc, 2, "2.6 Выводы по главе")
    add_body(doc, "Во второй главе разработана архитектура рекомендательной системы SciVise. Определены источники данных, признаки, компоненты гибридного скоринга, роль нейронного реранкера и механизм подмешивания публикаций из смежных направлений.")
    add_body(doc, "Предложенная архитектура соответствует целям MVP: она воспроизводима локально, не требует тяжелой инфраструктуры, остается объяснимой и позволяет постепенно усложнять алгоритмы. В следующей части отчета целесообразно описать реализацию пользовательского интерфейса и интеграцию модулей.")

    move_new_body_elements_before(doc, marker, original_len, section_properties)


def replace_chapter3(doc: Document, assets: dict[str, Path]) -> None:
    marker = remove_range_between(doc, "Frontend-разработка / модификация", "Заключение")
    body = doc._body._element
    original_len = len(list(body))
    section_properties = list(body)[-1]

    h1 = add_heading(doc, 1, "3 Frontend-разработка и модификация пользовательского интерфейса SciVise")
    h1.paragraph_format.page_break_before = True

    add_heading(doc, 2, "3.1 Разработка прототипа интерфейса")
    add_body(doc, "Пользовательский интерфейс SciVise разрабатывается как основная точка взаимодействия пользователя с рекомендательной системой. Через интерфейс пользователь просматривает ленту, открывает карточку статьи, сохраняет материалы, управляет профилем и изменяет параметры отображения.")
    add_body(doc, "Прототип ориентирован на исследовательский сценарий: пользователь должен быстро оценить публикацию, понять ее тему, сохранить материал и вернуться к ленте. Поэтому интерфейс построен вокруг карточек статей, тематических тегов, модальных окон и нижней навигации.")
    add_table(
        doc,
        "Таблица 6 — Основные модули клиентского приложения",
        [
            ["Модуль", "Назначение", "Роль в пользовательском сценарии"],
            ["HomeFeed", "отображение персональной ленты и тематических тегов", "первичный выбор публикаций"],
            ["ArticleCard", "краткое представление статьи", "быстрая оценка релевантности"],
            ["ArticleModal", "подробная карточка публикации", "чтение аннотации, комментарии, сохранение"],
            ["CollectionsFeed", "работа с подборками сохраненных статей", "организация материалов для дальнейшего изучения"],
            ["ProfileView", "профиль пользователя и статистика", "контроль личных данных и активности"],
            ["EditProfileModal", "редактирование профиля и настроек", "изменение языка, темы и персональных данных"],
            ["Navbar", "основная навигация приложения", "быстрый переход между лентой, коллекциями и профилем"],
        ],
    )
    add_figure(
        doc,
        assets["flow"],
        6.5,
        "Рисунок 10 — Логика пользовательских переходов в интерфейсе SciVise",
        "На рисунке 10 показаны основные маршруты пользователя: переход от гостевого режима к регистрации, настройке интересов, ленте, карточке статьи, коллекциям, профилю и настройкам приложения.",
    )
    add_figure(
        doc,
        assets["feed"],
        6.5,
        "Рисунок 11 — Главная лента научных публикаций",
        "На рисунке 11 представлен основной экран приложения. В ленте пользователь видит карточки публикаций, тематические теги и навигацию, что позволяет быстро перейти от общего просмотра к чтению конкретной статьи.",
    )

    add_heading(doc, 2, "3.2 Разработка логики переходов")
    add_body(doc, "Логика переходов построена вокруг двух режимов: гостевого и авторизованного. Гость может открыть приложение, просмотреть ленту и перейти к регистрации. Авторизованный пользователь получает доступ к коллекциям, профилю, настройкам языка и персонализированным действиям.")
    add_body(doc, "После регистрации пользователь проходит базовый сценарий настройки профиля. Он выбирает интересы, добавляет информацию о себе и затем попадает в ленту. Это позволяет рекомендательной системе сразу получить первичные тематические предпочтения.")
    add_body(doc, "Внутри ленты пользователь может открыть карточку статьи, сохранить публикацию в коллекцию, оставить комментарий или вернуться к списку. Такой сценарий соответствует типичной исследовательской работе: быстрый просмотр, углубленное чтение и сохранение значимых материалов.")
    add_figure(
        doc,
        assets["collections"],
        6.4,
        "Рисунок 12 — Экран коллекций сохраненных публикаций",
        "На рисунке 12 показан экран коллекций. Он нужен для группировки найденных публикаций по пользовательским подборкам и поддерживает повторное обращение к материалам при подготовке исследования.",
    )
    add_figure(
        doc,
        assets["profile"],
        6.4,
        "Рисунок 13 — Профиль пользователя и связанные разделы",
        "На рисунке 13 отображен профиль пользователя. В этом разделе собраны личные данные, сохраненные материалы и переходы к настройкам, которые влияют на дальнейшую персонализацию интерфейса.",
    )

    add_heading(doc, 2, "3.3 Технологическая реализация интерфейса")
    add_body(doc, "Клиентская часть реализована на React, TypeScript и Vite. React отвечает за компонентную структуру интерфейса, TypeScript снижает риск ошибок типов, Vite обеспечивает быструю сборку, а Tailwind CSS используется для адаптивной верстки и единообразных визуальных состояний.")
    add_body(doc, "Для связи с сервером используется отдельный слой API-сервисов. Такой подход отделяет интерфейсные компоненты от деталей HTTP-запросов и упрощает дальнейшее изменение backend-контрактов. Компоненты получают уже подготовленные данные и отвечают только за отображение и пользовательские действия.")
    add_body(doc, "Отдельное внимание уделено локализации. Интерфейс поддерживает русский и английский языки, а пользователь может выбрать язык в настройках профиля. Это важно, потому что язык системы и предпочтительный язык приложения могут различаться.")
    add_body(doc, "Мобильная адаптация реализована через нижнюю навигацию, компактные теги, устойчивые размеры карточек и модальные окна, которые не перекрывают основной контент. Такой подход позволяет использовать приложение на узких экранах без потери основных сценариев.")
    add_figure(
        doc,
        assets["edit_profile"],
        6.4,
        "Рисунок 14 — Редактирование профиля и пользовательских настроек",
        "На рисунке 14 показано окно редактирования профиля. Через него пользователь может изменить персональные данные, выбрать язык интерфейса и настроить отображение приложения под собственные предпочтения.",
    )
    add_figure(
        doc,
        assets["mobile"],
        3.3,
        "Рисунок 15 — Мобильное представление ленты публикаций",
        "На рисунке 15 показана мобильная версия ленты. Компактная навигация и адаптивные карточки позволяют сохранить основные сценарии приложения на узком экране.",
    )

    add_heading(doc, 2, "3.4 Выводы по главе")
    add_body(doc, "В третьей главе описана frontend-реализация SciVise. Были выделены основные модули интерфейса, пользовательские переходы, работа с лентой, коллекциями, профилем, настройками языка и мобильной версией приложения.")
    add_body(doc, "Разработанный интерфейс обеспечивает доступ к функциям рекомендательной системы и формирует пользовательские сигналы, необходимые для персонализации. Это связывает frontend-часть с общей архитектурой проекта и делает приложение пригодным для дальнейшей экспериментальной проверки.")

    move_new_body_elements_before(doc, marker, original_len, section_properties)


def replace_conclusion(doc: Document) -> None:
    marker = remove_range_between(doc, "Заключение", "Список используемых источников")
    body = doc._body._element
    original_len = len(list(body))
    section_properties = list(body)[-1]

    h1 = add_heading(doc, 1, "Заключение")
    h1.paragraph_format.page_break_before = True

    conclusion_paragraphs = [
        "В ходе преддипломной практики был проведен анализ, проектирование и описание прототипа информационной системы SciVise, предназначенной для поиска, сохранения и персонализированной рекомендации научных публикаций.",
        "Цель практики достигнута: сформирована проектная основа для выпускной квалификационной работы, определены требования к системе, описана архитектура приложения и подготовлено обоснование гибридного рекомендательного подхода.",
        "В первой главе была рассмотрена предметная область научного поиска. Были выделены основные пользователи системы, автоматизируемые процессы, вычислительные ресурсы, функциональные требования и состав работ по созданию MVP.",
        "Во второй главе была спроектирована архитектура рекомендательной системы. Отдельное внимание уделено данным, признакам, гибридному скорингу, нейронному реранкеру и механизму подмешивания публикаций из смежных научных направлений.",
        "В третьей главе была описана frontend-реализация пользовательского интерфейса. Рассмотрены лента публикаций, карточка статьи, коллекции, профиль, настройки языка, мобильная адаптация и связь интерфейса с пользовательскими сигналами.",
        "Практическая значимость выполненной работы заключается в том, что проект SciVise демонстрирует полный цикл работы с научными публикациями: от отображения материалов до накопления действий пользователя и формирования персонализированной выдачи.",
        "Полученные результаты могут быть использованы при дальнейшем выполнении ВКР. На следующем этапе целесообразно уточнить метрики качества рекомендаций, расширить набор публикаций и провести экспериментальное сравнение вариантов ранжирования.",
        "Таким образом, преддипломная практика позволила перейти от общего замысла проекта к оформленной архитектуре и рабочему прототипу, который может служить основой для дальнейшей разработки, тестирования и защиты выпускной квалификационной работы.",
    ]
    for text in conclusion_paragraphs:
        add_body(doc, text)

    move_new_body_elements_before(doc, marker, original_len, section_properties)


def update_toc(doc: Document) -> None:
    toc_texts = [
        "1 Анализ бизнес-процесса разработки информационной системы рекомендации научных публикаций SciVise",
        "1.1 Исходные данные для отчета",
        "1.2 Анализ структуры предметной области",
        "1.3 Анализ требований пользователей",
        "1.4 Анализ вычислительных ресурсов",
        "1.4.1 Нормативная и организационная база",
        "1.4.2 Структура информационно-коммуникационной среды",
        "1.4.2.1 Сетевая инфраструктура",
        "1.4.2.2 Аппаратное обеспечение",
        "1.4.2.3 Программное обеспечение",
        "1.5 Служебные поручения в период практики",
        "1.6 Техническое задание на разработку",
        "1.6.1 Общие сведения",
        "1.6.2 Назначение и цели создания системы",
        "1.6.3 Объекты автоматизации",
        "1.6.4 Требования к системе",
        "1.6.5 Состав и содержание работ",
        "1.6.6 Порядок разработки",
        "1.6.7 Порядок контроля и приемки",
        "1.7 Выводы по главе",
        "2 Проектирование архитектуры рекомендательной системы SciVise",
        "2.1 Разработка предложений по архитектуре рекомендательного модуля",
        "2.2 Архитектурная спецификация компонентов рекомендательной системы",
        "2.3 Разработка прототипа модели рекомендаций",
        "2.3.1 Формирование набора данных и признаков",
        "2.3.2 Выбор модели и инструментальных средств",
        "2.4 Обучение и расчет рекомендаций",
        "2.5 Разработка программного модуля рекомендательной системы",
        "2.6 Выводы по главе",
        "3 Frontend-разработка и модификация пользовательского интерфейса SciVise",
        "3.1 Разработка прототипа интерфейса",
        "3.2 Разработка логики переходов",
        "3.3 Технологическая реализация интерфейса",
        "3.4 Выводы по главе",
        "Заключение",
        "Список используемых источников и литературы",
    ]
    toc_paragraphs = [p for p in doc.paragraphs if p.style and p.style.name.lower().startswith("toc")]
    for p, text in zip(toc_paragraphs[3:], toc_texts):
        p.text = text
        set_paragraph_font(p, 12)


def build_document() -> None:
    ASSETS.mkdir(parents=True, exist_ok=True)
    process = ASSETS / "chapter1_business_process.png"
    architecture = ASSETS / "chapter1_architecture.png"
    chart = ASSETS / "chapter1_component_chart.png"
    title_page = ASSETS / "title_page_2026.png"
    pipeline = ASSETS / "chapter2_recommender_pipeline.png"
    scoring = ASSETS / "chapter2_scoring_chart.png"
    neural = ASSETS / "chapter2_neural_reranker.png"
    adjacent = ASSETS / "chapter2_adjacent_domains.png"
    frontend_flow = ASSETS / "chapter3_frontend_flow.png"
    screenshot = ROOT / "output/playwright/hybrid-recommender-feed.png"
    article_modal = ROOT / "output/playwright/07-article-modal.png"
    collections_screen = ROOT / "output/playwright/09-collections.png"
    profile_screen = ROOT / "output/playwright/10-profile.png"
    edit_profile_screen = ROOT / "output/playwright/11-edit-profile.png"
    mobile_screen = ROOT / "output/playwright/14-mobile-feed.png"
    if not screenshot.exists():
        screenshot = ROOT / "output/playwright/01-home.png"
    if not article_modal.exists():
        article_modal = screenshot
    if not collections_screen.exists():
        collections_screen = screenshot
    if not profile_screen.exists():
        profile_screen = screenshot
    if not edit_profile_screen.exists():
        edit_profile_screen = screenshot
    if not mobile_screen.exists():
        mobile_screen = screenshot
    make_process_diagram(process)
    make_architecture_diagram(architecture)
    make_component_chart(chart)
    make_recommender_pipeline(pipeline)
    make_scoring_chart(scoring)
    make_neural_reranker_diagram(neural)
    make_adjacent_domains_diagram(adjacent)
    make_frontend_flow_diagram(frontend_flow)

    shutil.copyfile(TEMPLATE, OUTPUT)
    doc = Document(OUTPUT)
    update_toc(doc)
    replace_intro(doc)
    marker = remove_range_between(doc, "Анализ бизнес-процесса / ов", "проектирование АРХИТЕКТУРЫ")
    body = doc._body._element
    original_len = len(list(body))
    section_properties = list(body)[-1]

    h1 = add_heading(doc, 1, "1 Анализ бизнес-процесса разработки информационной системы рекомендации научных публикаций SciVise")
    h1.paragraph_format.page_break_before = True

    add_heading(doc, 2, "1.1 Исходные данные для отчета")
    add_body(doc, "В период преддипломной практики рассматривался проект SciVise — веб-приложение для поиска, чтения, сохранения и персонализированной рекомендации научных публикаций. Проект связан с темой выпускной квалификационной работы, ориентированной на разработку гибридной рекомендательной системы научных публикаций с применением методов интеллектуального анализа данных и нейронного ранжирования.")
    add_body(doc, "Предметная область проекта — поддержка исследовательской и учебной деятельности студентов, преподавателей и молодых исследователей, которым необходимо быстро находить релевантные статьи, группировать их по коллекциям и получать новые публикации из смежных научных направлений. Основной автоматизируемый процесс включает получение публикаций из внешних источников, нормализацию метаданных, отображение карточек статей, сбор пользовательских сигналов и последующее ранжирование выдачи.")
    add_body(doc, "В качестве исходных материалов использованы индивидуальное задание на практику, шаблон отчета, материалы проекта, результаты ручного и автоматизированного тестирования, а также состояние локального Docker-окружения. При подготовке главы учитывались принятые проектные решения: исключение мессенджера, списка друзей, блокировки пользователей и чат-функций из текущего объема MVP; реакции оставлены только для комментариев к статьям.")

    add_figure(
        doc,
        process,
        6.6,
        "Рисунок 1 — Обобщенная схема бизнес-процесса работы с научными публикациями",
        "На рисунке 1 показан общий процесс работы пользователя с научными публикациями: поиск и импорт материалов, нормализация данных, просмотр ленты, фиксация пользовательских сигналов, расчет рекомендаций и сохранение результатов в коллекции.",
    )
    add_table(
        doc,
        "Таблица 1 — Исходные данные проекта и окружение разработки",
        [
            ["Параметр", "Значение"],
            ["Направление подготовки", "09.03.03 Прикладная информатика, профиль «Искусственный интеллект и анализ данных»"],
            ["Вид практики", "Преддипломная практика, 13.04.2026–10.05.2026"],
            ["Место прохождения", "Кафедра информационных систем"],
            ["Репозиторий проекта", "https://github.com/Skolkoff/Thinklist"],
            ["Локальное развертывание", "Frontend — http://localhost:5740, Backend API — http://localhost:8000, ML-service — http://localhost:8001"],
            ["Демонстрационные данные", "100 публикаций, 29 тематических тегов, 25 пользователей в локальной базе на момент обследования"],
        ],
    )

    add_heading(doc, 2, "1.2 Анализ структуры предметной области")
    add_body(doc, "В предметной области научного поиска пользователь обычно проходит несколько повторяющихся этапов: формулирует информационную потребность, просматривает выдачу, уточняет интересы, сохраняет значимые материалы и возвращается к ним в дальнейшем. Для учебного и исследовательского контекста важны не только точные совпадения по ключевым словам, но и возможность находить смежные публикации, которые расширяют научный обзор и помогают сформировать исследовательскую гипотезу.")
    add_body(doc, "Структура проекта ориентирована на поэтапную разработку прикладной информационной системы: сначала формируются базовые сценарии работы с публикациями, затем добавляются персональные действия пользователя, после чего подключается рекомендательный контур. Такой порядок позволяет отделить пользовательский интерфейс от серверной логики и от вычислительной части, отвечающей за ранжирование материалов.")
    add_body(doc, "Проектная структура соответствует разделению ответственности между клиентским приложением, API-сервером, ML-сервисом и базой данных. Такой подход упрощает развитие системы: интерфейс отвечает за пользовательские сценарии, Backend API — за авторизацию и доменную логику, ML-service — за расчет рекомендаций, PostgreSQL — за долговременное хранение данных и событий.")

    add_heading(doc, 2, "1.3 Анализ требований пользователей")
    add_body(doc, "В системе выделены три основные группы пользователей: гость, зарегистрированный пользователь и разработчик/администратор учебного стенда. Гость может ознакомиться с лентой и поиском без регистрации, зарегистрированный пользователь получает персонализацию, сохранение статей и работу с коллекциями, разработчик контролирует качество данных, миграции, контейнеры и тестовые сценарии.")
    add_table(
        doc,
        "Таблица 2 — Роли пользователей и ключевые требования",
        [
            ["Роль", "Основные потребности", "Функции системы"],
            ["Гость", "Быстро оценить содержание сервиса без создания аккаунта", "Просмотр ленты, поиск, открытие карточки статьи, переход к регистрации"],
            ["Зарегистрированный пользователь", "Получать персональные рекомендации и хранить полезные материалы", "Профиль, язык интерфейса, сохранение в коллекции, комментарии, реакции к комментариям, рекомендации"],
            ["Исследователь/студент", "Находить статьи по теме ВКР и смежным направлениям", "Тематические теги, поиск, гибридная выдача, баланс точности и новизны"],
            ["Разработчик учебного стенда", "Проверять работоспособность и воспроизводимость проекта", "Docker Compose, smoke-тесты, API-контракты, контроль контейнеров"],
        ],
    )
    add_body(doc, "Функциональные требования MVP включают регистрацию и вход, гостевой режим, настройку языка интерфейса, просмотр ленты, поиск публикаций, открытие карточки статьи, сохранение статьи в коллекцию, просмотр профиля, редактирование профиля, комментарии к статьям и получение рекомендаций. Нефункциональные требования включают воспроизводимое локальное развертывание, читаемый интерфейс на русском и английском языках, устойчивость к пустым данным, проверяемые API-контракты и отсутствие ошибок в основных пользовательских сценариях.")

    add_heading(doc, 2, "1.4 Анализ вычислительных ресурсов")
    add_heading(doc, 3, "1.4.1 Нормативная и организационная база")
    add_body(doc, "При описании технического задания использован подход, близкий к структуре ГОСТ 34.602-89: общие сведения, назначение системы, объекты автоматизации, требования, состав работ, порядок контроля и требования к документации. Для учебного проекта это позволяет формализовать результат практики и связать разработку приложения с ожидаемыми профессиональными компетенциями.")
    add_heading(doc, 3, "1.4.2 Структура информационно-коммуникационной среды")
    add_body(doc, "Среда проекта построена как набор контейнеров Docker Compose. На момент обследования в локальном окружении подняты frontend-контейнер, Backend API, ML-service и PostgreSQL. Такой состав достаточен для демонстрации полного цикла: пользователь открывает веб-интерфейс, API обрабатывает запросы и авторизацию, база данных хранит публикации и события, ML-service рассчитывает персонализированные рекомендации.")
    add_figure(
        doc,
        architecture,
        6.6,
        "Рисунок 2 — Архитектура информационной системы SciVise",
        "На рисунке 2 представлено разделение системы на frontend, Backend API, ML-service и PostgreSQL. Схема показывает, какие компоненты отвечают за интерфейс, бизнес-логику, вычисление рекомендаций и хранение данных.",
    )
    add_heading(doc, 4, "1.4.2.1 Сетевая инфраструктура")
    add_body(doc, "В локальном стенде используется внутренняя Docker-сеть scivise-network. Внешнему пользователю доступны порты 5740 для веб-интерфейса, 8000 для Backend API, 8001 для ML-service и 5432 для PostgreSQL. Для production-развертывания эти порты должны быть закрыты reverse proxy и политиками доступа, однако для преддипломной практики локальная конфигурация обеспечивает прозрачную проверку всех компонентов.")
    add_heading(doc, 4, "1.4.2.2 Аппаратное обеспечение")
    add_body(doc, "Для разработки и демонстрации достаточно рабочей станции разработчика с поддержкой Docker Desktop. Текущий объем данных невелик: демонстрационный набор содержит 100 публикаций, поэтому обучение и расчет признаков возможны на CPU без отдельного GPU. При увеличении корпуса публикаций потребуется выделить ресурсы для регулярного обновления эмбеддингов и хранения векторных представлений.")
    add_heading(doc, 4, "1.4.2.3 Программное обеспечение")
    add_body(doc, "Клиентская часть реализована на React 19, TypeScript, Vite и Tailwind CSS. Backend API написан на FastAPI с использованием SQLAlchemy и Alembic. ML-service также использует FastAPI и содержит гибридный рекомендательный пайплайн: TF-IDF, понижение размерности TruncatedSVD, нейронный реранкер на MLPRegressor, факторы свежести, интересов, поведения и разнообразия. Для хранения данных используется PostgreSQL 15.")
    add_figure(
        doc,
        screenshot,
        6.6,
        "Рисунок 3 — Скриншот ленты рекомендаций в веб-приложении",
        "На рисунке 3 показан основной экран приложения с лентой публикаций. Пользователь видит карточки статей, тематические фильтры и нижнюю навигацию, что позволяет быстро переходить между чтением, коллекциями и профилем.",
    )
    add_figure(
        doc,
        chart,
        6.3,
        "Рисунок 4 — Распределение элементов MVP по ключевым подсистемам",
        "На рисунке 4 обобщен состав MVP по основным подсистемам. Диаграмма показывает, что проект включает интерфейсные модули, API-группы, сущности базы данных, факторы ML-скоринга и сценарии проверки работоспособности.",
    )

    add_heading(doc, 2, "1.5 Служебные поручения в период практики")
    add_body(doc, "В ходе практики выполнялись работы, направленные на доведение проекта до состояния MVP, пригодного для демонстрации и дальнейшего развития. Основной акцент был сделан на стабилизации пользовательских сценариев, сокращении лишнего функционального объема и формализации рекомендательной системы.")
    add_numbered(
        doc,
        [
            "Проведен первичный аудит приложения, контейнеров, API и пользовательских сценариев.",
            "Проведена проверка ключевых пользовательских сценариев и исправлены выявленные проблемы.",
            "Из текущего объема MVP исключены мессенджер, друзья, блокировка пользователей и операции очистки чатов.",
            "Доработана карточка статьи, комментарии, коллекции, профиль пользователя и настройки языка интерфейса.",
            "Сформирован минимально достаточный стек гибридной рекомендательной системы для темы ВКР.",
            "Проверена работоспособность фронтенда, backend API, ML-service и контейнерной сборки.",
        ],
    )

    add_heading(doc, 2, "1.6 Техническое задание на разработку")
    add_heading(doc, 3, "1.6.1 Общие сведения")
    add_body(doc, "Наименование разрабатываемой системы: SciVise — информационная система поиска, хранения и рекомендации научных публикаций. Основание для разработки: выполнение преддипломной практики и подготовка выпускной квалификационной работы по направлению искусственного интеллекта и анализа данных.")
    add_heading(doc, 3, "1.6.2 Назначение и цели создания системы")
    add_body(doc, "Назначение системы заключается в поддержке исследовательской работы пользователя при выборе и изучении научных публикаций. Цель разработки — создать воспроизводимый MVP, который демонстрирует полный цикл персонализации: загрузка публикаций, интерфейс взаимодействия, сбор пользовательских сигналов и гибридное ранжирование материалов.")
    add_heading(doc, 3, "1.6.3 Объекты автоматизации")
    add_body(doc, "Объектами автоматизации являются процессы поиска научных публикаций, просмотра карточек статей, сохранения материалов, ведения пользовательских коллекций, фиксации действий пользователя и формирования персонализированной выдачи.")
    add_heading(doc, 3, "1.6.4 Требования к системе")
    add_body(doc, "Система должна обеспечивать гостевой и авторизованный режимы работы, регистрацию и вход, настройку языка интерфейса, просмотр ленты публикаций, поиск, открытие карточки статьи, сохранение материалов в коллекции, комментирование статьи, отображение профиля и получение рекомендаций. Интерфейс должен поддерживать русский и английский языки, адаптироваться к мобильной ширине экрана и корректно отображать короткие тематические теги.")
    add_body(doc, "Рекомендательная подсистема должна использовать гибридный подход: контентные признаки публикации, интересы пользователя, поведенческие события, фактор свежести, нейронный реранкер и механизм разнообразия. Для научного контекста важна возможность подмешивать статьи из смежных направлений, чтобы пользователь не замыкался только на узкой теме.")
    add_heading(doc, 3, "1.6.5 Состав и содержание работ")
    add_bullets(
        doc,
        [
            "подготовка схемы данных и миграций PostgreSQL;",
            "реализация Backend API для пользователей, публикаций, коллекций, комментариев и рекомендаций;",
            "создание ML-service для расчета гибридного скоринга;",
            "разработка клиентского интерфейса на React и TypeScript;",
            "локализация интерфейса на русский и английский языки;",
            "подготовка набора публикаций и проверка основных пользовательских сценариев.",
        ],
    )
    add_heading(doc, 3, "1.6.6 Порядок разработки")
    add_body(doc, "Разработка выполняется итерационно. На первом этапе формируются базовые сценарии и схема данных, на втором — пользовательский интерфейс и API-контракты, на третьем — рекомендательная система и набор публикаций, на четвертом — стабилизация, тестирование, контейнерная сборка и подготовка материалов к защите.")
    add_heading(doc, 3, "1.6.7 Порядок контроля и приемки")
    add_body(doc, "Контроль осуществляется через unit- и contract-тесты, smoke-тесты API, проверку сборки frontend, проверку Docker Compose и ручные сценарии в браузере. Для приемки MVP важны отсутствие ошибок консоли в основных сценариях, корректное отображение интерфейса на мобильной ширине, прохождение тестов рекомендательной подсистемы и воспроизводимый запуск контейнеров.")
    add_heading(doc, 3, "1.6.8 Подготовка объекта автоматизации к вводу в действие")
    add_body(doc, "Для демонстрационного ввода необходимо выполнить миграции базы данных, загрузить демо-публикации, запустить контейнеры Docker Compose, проверить health endpoints Backend API и ML-service, открыть веб-интерфейс на порту 5740 и пройти базовые сценарии гостя и авторизованного пользователя.")
    add_heading(doc, 3, "1.6.9 Требования к документированию")
    add_body(doc, "Документация должна включать описание архитектуры, стек технологий, структуру рекомендательного пайплайна, результаты тестирования, принятые решения по функциональному объему и инструкции по запуску системы.")
    add_heading(doc, 3, "1.6.10 Источники разработки")
    add_body(doc, "Источниками разработки являются материалы индивидуального задания, учебные требования к преддипломной практике, репозиторий проекта SciVise, документация используемых технологий и локальные результаты тестирования.")

    add_heading(doc, 2, "1.7 Выводы по главе")
    add_body(doc, "В первой главе проведен анализ предметной области и текущего состояния проекта SciVise. Установлено, что проект находится на стадии MVP: основные пользовательские сценарии реализованы, сервисы контейнеризованы, рекомендательная система имеет минимально достаточный гибридный пайплайн, а дальнейшая работа должна быть сосредоточена на доведении качества данных, метрик рекомендаций и итоговых материалов.")
    add_table(
        doc,
        "Таблица 3 — Итоги анализа первой главы",
        [
            ["Направление анализа", "Результат", "Дальнейшие действия"],
            ["Предметная область", "Определен процесс поиска, отбора и персонализации научных публикаций", "Уточнить формальные критерии качества рекомендаций"],
            ["Функциональный объем", "MVP очищен от лишних социальных функций и сфокусирован на статьях", "Закрепить скоуп в последующих главах отчета"],
            ["Архитектура", "Выделены frontend, Backend API, ML-service и PostgreSQL", "Описать реализацию модулей во второй главе"],
            ["Рекомендации", "Сформирован гибридный стек с нейронным реранкером", "Подготовить экспериментальную часть и метрики"],
            ["Тестирование", "Зафиксированы smoke-, contract- и browser-проверки", "Повторить полный прогон перед финальной сдачей проекта"],
        ],
    )

    move_new_body_elements_before(doc, marker, original_len, section_properties)
    replace_chapter2(
        doc,
        {
            "pipeline": pipeline,
            "scoring": scoring,
            "neural": neural,
            "adjacent": adjacent,
            "article_modal": article_modal,
        },
    )
    replace_chapter3(
        doc,
        {
            "flow": frontend_flow,
            "feed": screenshot,
            "collections": collections_screen,
            "profile": profile_screen,
            "edit_profile": edit_profile_screen,
            "mobile": mobile_screen,
        },
    )
    replace_conclusion(doc)

    doc.save(OUTPUT)
    with zipfile.ZipFile(OUTPUT, "r") as z:
        title_bytes = z.read("word/media/image1.png")
    make_title_page_image(title_bytes, title_page)
    replace_docx_media(OUTPUT, "word/media/image1.png", title_page)


if __name__ == "__main__":
    build_document()
    print(OUTPUT)
